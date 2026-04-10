"""
Excel Formatter - Professional styling for deliverables
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any
import pandas as pd


class ExcelFormatter:
    """Professional Excel formatting for engineering deliverables"""
    
    # Color scheme
    HEADER_COLOR = "1F4E78"  # Dark blue
    ALT_ROW_COLOR = "F2F2F2"  # Light gray
    ALARM_HIGH_COLOR = "FFC7CE"  # Light red
    ALARM_LOW_COLOR = "FFEB9C"  # Light yellow
    INTERLOCK_COLOR = "C6EFCE"  # Light green
    
    def __init__(self):
        self.workbook = Workbook()
        # Remove default sheet
        if 'Sheet' in self.workbook.sheetnames:
            self.workbook.remove(self.workbook['Sheet'])
    
    def create_formatted_sheet(
        self,
        sheet_name: str,
        data: pd.DataFrame,
        title: str,
        subtitle: str = None
    ) -> None:
        """Create a professionally formatted worksheet"""
        
        ws = self.workbook.create_sheet(sheet_name)
        
        # Add title
        ws.merge_cells('A1:' + get_column_letter(len(data.columns)) + '1')
        title_cell = ws['A1']
        title_cell.value = title
        title_cell.font = Font(size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color=self.HEADER_COLOR, 
                                      end_color=self.HEADER_COLOR, 
                                      fill_type="solid")
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # Add subtitle if provided
        start_row = 2
        if subtitle:
            ws.merge_cells(f'A2:{get_column_letter(len(data.columns))}2')
            subtitle_cell = ws['A2']
            subtitle_cell.value = subtitle
            subtitle_cell.font = Font(size=11, italic=True)
            subtitle_cell.alignment = Alignment(horizontal='center')
            ws.row_dimensions[2].height = 20
            start_row = 3
        
        # Add headers
        header_row = start_row + 1
        for col_idx, column in enumerate(data.columns, 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.value = column
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color=self.HEADER_COLOR,
                                   end_color=self.HEADER_COLOR,
                                   fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self._get_border()
        
        ws.row_dimensions[header_row].height = 25
        
        # Add data rows
        for row_idx, row_data in enumerate(data.itertuples(index=False), header_row + 1):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = value
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.border = self._get_border()
                
                # Alternate row colors
                if row_idx % 2 == 0:
                    cell.fill = PatternFill(start_color=self.ALT_ROW_COLOR,
                                          end_color=self.ALT_ROW_COLOR,
                                          fill_type="solid")
                
                # Special formatting for alarm/interlock indicators
                if isinstance(value, str):
                    if 'HIGH' in value.upper() or 'HH' in value.upper():
                        cell.fill = PatternFill(start_color=self.ALARM_HIGH_COLOR,
                                              end_color=self.ALARM_HIGH_COLOR,
                                              fill_type="solid")
                    elif 'LOW' in value.upper() or 'LL' in value.upper():
                        cell.fill = PatternFill(start_color=self.ALARM_LOW_COLOR,
                                              end_color=self.ALARM_LOW_COLOR,
                                              fill_type="solid")
                    elif 'INTERLOCK' in value.upper() or 'TRIP' in value.upper():
                        cell.fill = PatternFill(start_color=self.INTERLOCK_COLOR,
                                              end_color=self.INTERLOCK_COLOR,
                                              fill_type="solid")
        
        # Auto-adjust column widths
        self._auto_adjust_columns(ws, data)
        
        # Freeze header row
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    
    def _get_border(self) -> Border:
        """Get standard border style"""
        thin = Side(border_style="thin", color="000000")
        return Border(left=thin, right=thin, top=thin, bottom=thin)
    
    def _auto_adjust_columns(self, ws, data: pd.DataFrame):
        """Auto-adjust column widths based on content"""
        for col_idx, column in enumerate(data.columns, 1):
            max_length = len(str(column))
            for row in data[column]:
                if row is not None:
                    max_length = max(max_length, len(str(row)))
            
            # Set width with some padding
            adjusted_width = min(max_length + 3, 50)  # Max 50 chars
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
    
    def save(self, filepath: str):
        """Save the workbook"""
        self.workbook.save(filepath)
